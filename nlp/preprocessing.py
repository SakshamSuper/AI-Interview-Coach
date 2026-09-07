import re
import os
from typing import Dict, List, Optional
import pypdf
import docx


def clean_text(text: str) -> str:
    """Cleans and standardizes raw text extracted from documents."""
    if not text:
        return ""
    # Normalize line breaks
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    # Remove null bytes and non-printable control characters
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text)
    # Collapse multiple spaces but preserve single newlines
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.split("\n")]
    # Remove excessive blank lines (more than 2 consecutive)
    cleaned = re.sub(r"\n{3,}", "\n\n", "\n".join(lines))
    return cleaned.strip()


def extract_text_from_file(file_path: str) -> str:
    """Extracts raw text from PDF, DOCX, or TXT files."""
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")

    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext in [".docx", ".doc"]:
        return extract_text_from_docx(file_path)
    elif ext in [".txt", ".md"]:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return clean_text(f.read())
    else:
        # Fallback to UTF-8 text read
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return clean_text(f.read())


def extract_text_from_pdf(pdf_path: str) -> str:
    """Extracts text from PDF using pypdf."""
    text_content = []
    with open(pdf_path, "rb") as f:
        reader = pypdf.PdfReader(f)
        for page_idx, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text:
                text_content.append(page_text)
    return clean_text("\n".join(text_content))


def extract_text_from_docx(docx_path: str) -> str:
    """Extracts text from DOCX using python-docx."""
    doc = docx.Document(docx_path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                paragraphs.append(row_text)
    return clean_text("\n".join(paragraphs))


# Known section header aliases for resume parsing
SECTION_ALIASES: Dict[str, List[str]] = {
    "summary": ["summary", "professional summary", "about me", "profile", "objective", "career objective"],
    "skills": ["skills", "technical skills", "core competencies", "skills & tools", "technologies", "expertise"],
    "experience": ["experience", "work experience", "professional experience", "employment history", "work history"],
    "education": ["education", "academic background", "educational qualifications", "academics"],
    "projects": ["projects", "personal projects", "academic projects", "key projects"],
    "certifications": ["certifications", "licenses & certifications", "courses & certifications", "accreditations"],
    "achievements": ["achievements", "honors", "awards", "accomplishments", "publications"]
}


def detect_sections(text: str) -> Dict[str, str]:
    """
    Partitions resume text into semantic sections based on standard section headers.
    """
    lines = text.split("\n")
    sections: Dict[str, List[str]] = {"header": []}
    current_section = "header"

    # Regex to match a potential section header line
    header_pattern = re.compile(r"^[A-Z0-9\s,&/|•\-_:]{2,40}$")

    for line in lines:
        stripped = line.strip().lower()
        # Clean punctuation for matching
        clean_line = re.sub(r"[:\-_|•]", "", stripped).strip()

        matched_section = None
        for sec_name, aliases in SECTION_ALIASES.items():
            if clean_line in aliases:
                matched_section = sec_name
                break

        if matched_section and len(line.strip().split()) <= 5:
            current_section = matched_section
            if current_section not in sections:
                sections[current_section] = []
        else:
            if current_section not in sections:
                sections[current_section] = []
            sections[current_section].append(line)

    return {sec: "\n".join(lines).strip() for sec, lines in sections.items() if "\n".join(lines).strip()}
